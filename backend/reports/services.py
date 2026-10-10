import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from io import BytesIO
from django.db.models import Sum, Count, Avg, F
from django.utils import timezone
from vehicles.models import Vehicle, VehicleType, OdometerReading, VehicleDocument
from drivers.models import Driver
from work_permits.models import WorkPermit
from maintenance.models import MaintenanceRecord, MaintenanceType
from parts.models import VehiclePart, PartCatalog
from fuel.models import FuelRecord


class ReportGeneratorService:
    @classmethod
    def generate_report_data(cls, report_type, filters, user=None):
        date_from = filters.get('date_from')
        date_to = filters.get('date_to')
        vehicle_id = filters.get('vehicle')
        driver_id = filters.get('driver')
        status_param = filters.get('status')
        type_id = filters.get('type')
        department = filters.get('department')

        data = []
        columns = []
        summary = {}
        title = ""

        # 1. Vehicle Report
        if report_type == 'VEHICLES':
            title = "تقرير أسطول المركبات الشامل"
            columns = [
                {'key': 'fleet_number', 'label': 'الرقم الداخلي'},
                {'key': 'plate_number', 'label': 'رقم اللوحة'},
                {'key': 'type_name', 'label': 'النوع'},
                {'key': 'make_model', 'label': 'الماركة والموديل'},
                {'key': 'year', 'label': 'سنة الصنع'},
                {'key': 'current_odometer', 'label': 'العداد الحالي (كم)'},
                {'key': 'status', 'label': 'الحالة'},
                {'key': 'driver_name', 'label': 'السائق'},
                {'key': 'department', 'label': 'القسم'},
            ]
            qs = Vehicle.objects.filter(is_deleted=False).select_related('vehicle_type', 'assigned_driver')
            if vehicle_id:
                qs = qs.filter(id=vehicle_id)
            if status_param:
                qs = qs.filter(status=status_param)
            if department:
                qs = qs.filter(department__icontains=department)

            for v in qs:
                data.append({
                    'fleet_number': v.internal_fleet_number,
                    'plate_number': v.plate_number,
                    'type_name': v.vehicle_type.name,
                    'make_model': f"{v.make} {v.model}",
                    'year': v.year,
                    'current_odometer': f"{v.current_odometer:,}",
                    'status': v.get_status_display(),
                    'driver_name': v.assigned_driver.full_name if v.assigned_driver else 'غير معين',
                    'department': v.department or 'عام',
                })
            summary = {'إجمالي المركبات': len(data)}

        # 2. Vehicle Usage Report
        elif report_type == 'VEHICLE_USAGE':
            title = "تقرير تشغيل واستخدام المركبات"
            columns = [
                {'key': 'fleet_number', 'label': 'الرقم الداخلي'},
                {'key': 'plate_number', 'label': 'رقم اللوحة'},
                {'key': 'total_trips', 'label': 'عدد الرحلات'},
                {'key': 'total_distance', 'label': 'إجمالي المسافة (كم)'},
                {'key': 'total_fuel_liters', 'label': 'استهلاك الوقود (لتر)'},
                {'key': 'total_fuel_cost', 'label': 'تكلفة الوقود (ر.س)'},
            ]
            qs = Vehicle.objects.filter(is_deleted=False)
            if vehicle_id:
                qs = qs.filter(id=vehicle_id)
            total_dist_fleet = 0
            total_fuel_fleet = 0

            for v in qs:
                permits = v.work_permits.filter(status=WorkPermit.Status.COMPLETED)
                if date_from:
                    permits = permits.filter(departure_date__gte=date_from)
                if date_to:
                    permits = permits.filter(departure_date__lte=date_to)
                
                dist = permits.aggregate(total=Sum('distance_traveled'))['total'] or 0
                trips_count = permits.count()

                fuel_recs = v.fuel_records.filter(is_deleted=False)
                if date_from:
                    fuel_recs = fuel_recs.filter(date__gte=date_from)
                if date_to:
                    fuel_recs = fuel_recs.filter(date__lte=date_to)
                
                fuel_agg = fuel_recs.aggregate(liters=Sum('quantity'), cost=Sum('total_cost'))
                liters = fuel_agg['liters'] or 0
                cost = fuel_agg['cost'] or 0

                total_dist_fleet += dist
                total_fuel_fleet += cost

                data.append({
                    'fleet_number': v.internal_fleet_number,
                    'plate_number': v.plate_number,
                    'total_trips': trips_count,
                    'total_distance': f"{dist:,}",
                    'total_fuel_liters': f"{liters:,.1f}",
                    'total_fuel_cost': f"{cost:,.2f}",
                })
            summary = {
                'إجمالي المسافة المقطوعة بالأسطول': f"{total_dist_fleet:,} كم",
                'إجمالي تكلفة الوقود': f"{total_fuel_fleet:,.2f} ر.س"
            }

        # 3. Work Permit Report
        elif report_type == 'WORK_PERMITS':
            title = "تقرير أذونات التشغيل والرحلات"
            columns = [
                {'key': 'permit_number', 'label': 'رقم الإذن'},
                {'key': 'vehicle', 'label': 'المركبة'},
                {'key': 'driver', 'label': 'السائق'},
                {'key': 'departure_date', 'label': 'تاريخ الانطلاق'},
                {'key': 'destination', 'label': 'الوجهة'},
                {'key': 'purpose', 'label': 'الغرض'},
                {'key': 'distance', 'label': 'المسافة (كم)'},
                {'key': 'status', 'label': 'الحالة'},
            ]
            qs = WorkPermit.objects.filter(is_deleted=False).select_related('vehicle', 'driver')
            if date_from:
                qs = qs.filter(departure_date__gte=date_from)
            if date_to:
                qs = qs.filter(departure_date__lte=date_to)
            if vehicle_id:
                qs = qs.filter(vehicle_id=vehicle_id)
            if driver_id:
                qs = qs.filter(driver_id=driver_id)
            if status_param:
                qs = qs.filter(status=status_param)

            total_dist = 0
            for wp in qs:
                total_dist += wp.distance_traveled
                data.append({
                    'permit_number': wp.permit_number,
                    'vehicle': f"{wp.vehicle.internal_fleet_number} - {wp.vehicle.plate_number}",
                    'driver': wp.driver.full_name,
                    'departure_date': str(wp.departure_date),
                    'destination': wp.destination,
                    'purpose': wp.purpose,
                    'distance': f"{wp.distance_traveled:,}",
                    'status': wp.get_status_display(),
                })
            summary = {
                'عدد الأذونات': len(data),
                'إجمالي المسافة المقطوعة': f"{total_dist:,} كم"
            }

        # 4. Trip Report (Completed Trips)
        elif report_type == 'TRIPS':
            title = "تقرير الرحلات المكتملة"
            columns = [
                {'key': 'permit_number', 'label': 'رقم الرحلة'},
                {'key': 'vehicle', 'label': 'المركبة'},
                {'key': 'driver', 'label': 'السائق'},
                {'key': 'departure', 'label': 'الانطلاق'},
                {'key': 'return', 'label': 'العودة'},
                {'key': 'starting_odometer', 'label': 'عداد البداية'},
                {'key': 'ending_odometer', 'label': 'عداد العودة'},
                {'key': 'distance', 'label': 'المسافة (كم)'},
            ]
            qs = WorkPermit.objects.filter(is_deleted=False, status=WorkPermit.Status.COMPLETED).select_related('vehicle', 'driver')
            if date_from:
                qs = qs.filter(departure_date__gte=date_from)
            if date_to:
                qs = qs.filter(departure_date__lte=date_to)
            if vehicle_id:
                qs = qs.filter(vehicle_id=vehicle_id)
            if driver_id:
                qs = qs.filter(driver_id=driver_id)

            total_km = 0
            for wp in qs:
                total_km += wp.distance_traveled
                data.append({
                    'permit_number': wp.permit_number,
                    'vehicle': f"{wp.vehicle.internal_fleet_number} ({wp.vehicle.plate_number})",
                    'driver': wp.driver.full_name,
                    'departure': f"{wp.departure_date} {wp.departure_time or ''}".strip(),
                    'return': f"{wp.return_date or ''} {wp.return_time or ''}".strip(),
                    'starting_odometer': f"{wp.starting_odometer:,}" if wp.starting_odometer else '-',
                    'ending_odometer': f"{wp.ending_odometer:,}" if wp.ending_odometer else '-',
                    'distance': f"{wp.distance_traveled:,}",
                })
            summary = {'إجمالي الرحلات المكتملة': len(data), 'إجمالي المسافة': f"{total_km:,} كم"}

        # 5. Maintenance Report
        elif report_type == 'MAINTENANCE':
            title = "تقرير عمليات الصيانة والإصلاح"
            columns = [
                {'key': 'vehicle', 'label': 'المركبة'},
                {'key': 'maintenance_type', 'label': 'نوع الصيانة'},
                {'key': 'date', 'label': 'التاريخ'},
                {'key': 'odometer', 'label': 'العداد (كم)'},
                {'key': 'cost', 'label': 'التكلفة (ر.س)'},
                {'key': 'service_provider', 'label': 'مزود الخدمة'},
                {'key': 'status', 'label': 'الحالة'},
            ]
            qs = MaintenanceRecord.objects.filter(is_deleted=False).select_related('vehicle', 'maintenance_type')
            if date_from:
                qs = qs.filter(maintenance_date__gte=date_from)
            if date_to:
                qs = qs.filter(maintenance_date__lte=date_to)
            if vehicle_id:
                qs = qs.filter(vehicle_id=vehicle_id)
            if type_id:
                qs = qs.filter(maintenance_type_id=type_id)

            total_cost = 0
            for m in qs:
                total_cost += float(m.cost)
                data.append({
                    'vehicle': f"{m.vehicle.internal_fleet_number} ({m.vehicle.plate_number})",
                    'maintenance_type': m.maintenance_type.name,
                    'date': str(m.maintenance_date),
                    'odometer': f"{m.odometer:,}",
                    'cost': f"{m.cost:,.2f}",
                    'service_provider': m.service_provider or '-',
                    'status': m.get_status_display(),
                })
            summary = {'إجمالي العمليات': len(data), 'إجمالي تكلفة الصيانة': f"{total_cost:,.2f} ر.س"}

        # 6. Maintenance Cost Report
        elif report_type == 'MAINTENANCE_COST':
            title = "تقرير تحليل تكاليف الصيانة حسب المركبات"
            columns = [
                {'key': 'fleet_number', 'label': 'الرقم الداخلي'},
                {'key': 'plate_number', 'label': 'رقم اللوحة'},
                {'key': 'type_name', 'label': 'نوع المركبة'},
                {'key': 'records_count', 'label': 'عدد العمليات'},
                {'key': 'total_cost', 'label': 'إجمالي التكلفة (ر.س)'},
                {'key': 'avg_cost', 'label': 'متوسط تكلفة العملية'},
            ]
            qs = Vehicle.objects.filter(is_deleted=False)
            if vehicle_id:
                qs = qs.filter(id=vehicle_id)

            overall_cost = 0
            for v in qs:
                m_qs = v.maintenance_records.filter(is_deleted=False)
                if date_from:
                    m_qs = m_qs.filter(maintenance_date__gte=date_from)
                if date_to:
                    m_qs = m_qs.filter(maintenance_date__lte=date_to)
                
                agg = m_qs.aggregate(total=Sum('cost'), count=Count('id'), avg=Avg('cost'))
                c = agg['count'] or 0
                tot = float(agg['total'] or 0)
                avg = float(agg['avg'] or 0)
                overall_cost += tot

                if c > 0 or not date_from:
                    data.append({
                        'fleet_number': v.internal_fleet_number,
                        'plate_number': v.plate_number,
                        'type_name': v.vehicle_type.name,
                        'records_count': c,
                        'total_cost': f"{tot:,.2f}",
                        'avg_cost': f"{avg:,.2f}",
                    })
            summary = {'إجمالي تكاليف الصيانة لكافة المركبات': f"{overall_cost:,.2f} ر.س"}

        # 7. Parts Replacement Report
        elif report_type == 'PARTS_REPLACEMENT':
            title = "تقرير أداء واستبدال قطع الغيار"
            columns = [
                {'key': 'vehicle', 'label': 'المركبة'},
                {'key': 'part_name', 'label': 'اسم القطعة'},
                {'key': 'installation_date', 'label': 'تاريخ التركيب'},
                {'key': 'installation_odometer', 'label': 'عداد التركيب'},
                {'key': 'interval_km', 'label': 'فاصل الاستبدال (كم)'},
                {'key': 'next_km', 'label': 'العداد المستهدف'},
                {'key': 'remaining_km', 'label': 'المتبقي (كم)'},
                {'key': 'status', 'label': 'الحالة'},
            ]
            qs = VehiclePart.objects.all().select_related('vehicle', 'part')
            if vehicle_id:
                qs = qs.filter(vehicle_id=vehicle_id)
            if status_param:
                qs = qs.filter(current_status=status_param)

            for p in qs:
                rem_km = p.next_replacement_km - p.vehicle.current_odometer
                data.append({
                    'vehicle': f"{p.vehicle.internal_fleet_number} ({p.vehicle.plate_number})",
                    'part_name': p.part.name,
                    'installation_date': str(p.installation_date),
                    'installation_odometer': f"{p.installation_odometer:,}",
                    'interval_km': f"{p.interval_km:,}",
                    'next_km': f"{p.next_replacement_km:,}",
                    'remaining_km': f"{rem_km:,}",
                    'status': p.get_current_status_display(),
                })
            summary = {'إجمالي القطع المراقبة': len(data)}

        # 8. Fuel Report
        elif report_type == 'FUEL':
            title = "تقرير سجلات تعبئة الوقود"
            columns = [
                {'key': 'vehicle', 'label': 'المركبة'},
                {'key': 'date', 'label': 'التاريخ'},
                {'key': 'odometer', 'label': 'العداد (كم)'},
                {'key': 'fuel_type', 'label': 'نوع الوقود'},
                {'key': 'quantity', 'label': 'الكمية (لتر)'},
                {'key': 'price_per_unit', 'label': 'سعر اللتر'},
                {'key': 'total_cost', 'label': 'التكلفة الإجمالية (ر.س)'},
                {'key': 'station', 'label': 'المحطة'},
                {'key': 'driver', 'label': 'السائق'},
            ]
            qs = FuelRecord.objects.filter(is_deleted=False).select_related('vehicle', 'driver')
            if date_from:
                qs = qs.filter(date__gte=date_from)
            if date_to:
                qs = qs.filter(date__lte=date_to)
            if vehicle_id:
                qs = qs.filter(vehicle_id=vehicle_id)
            if driver_id:
                qs = qs.filter(driver_id=driver_id)

            tot_liters = 0
            tot_cost = 0
            for f in qs:
                tot_liters += float(f.quantity)
                tot_cost += float(f.total_cost)
                data.append({
                    'vehicle': f"{f.vehicle.internal_fleet_number} ({f.vehicle.plate_number})",
                    'date': str(f.date),
                    'odometer': f"{f.odometer:,}",
                    'fuel_type': f.get_fuel_type_display(),
                    'quantity': f"{f.quantity:,.1f}",
                    'price_per_unit': f"{f.price_per_unit:,.2f}",
                    'total_cost': f"{f.total_cost:,.2f}",
                    'station': f.station or '-',
                    'driver': f.driver.full_name if f.driver else '-',
                })
            summary = {
                'إجمالي كمية الوقود': f"{tot_liters:,.1f} لتر",
                'إجمالي تكلفة الوقود': f"{tot_cost:,.2f} ر.س"
            }

        # 9. Fuel Cost Report
        elif report_type == 'FUEL_COST':
            title = "تقرير تحليل تكاليف الوقود حسب المركبات"
            columns = [
                {'key': 'fleet_number', 'label': 'الرقم الداخلي'},
                {'key': 'plate_number', 'label': 'رقم اللوحة'},
                {'key': 'fuel_type', 'label': 'نوع الوقود'},
                {'key': 'refuel_count', 'label': 'مرات التعبئة'},
                {'key': 'total_liters', 'label': 'إجمالي اللترات'},
                {'key': 'total_cost', 'label': 'إجمالي التكلفة (ر.س)'},
            ]
            qs = Vehicle.objects.filter(is_deleted=False)
            if vehicle_id:
                qs = qs.filter(id=vehicle_id)

            total_lit = 0
            total_money = 0
            for v in qs:
                f_qs = v.fuel_records.filter(is_deleted=False)
                if date_from:
                    f_qs = f_qs.filter(date__gte=date_from)
                if date_to:
                    f_qs = f_qs.filter(date__lte=date_to)
                
                agg = f_qs.aggregate(liters=Sum('quantity'), cost=Sum('total_cost'), cnt=Count('id'))
                cnt = agg['cnt'] or 0
                lit = float(agg['liters'] or 0)
                cst = float(agg['cost'] or 0)
                total_lit += lit
                total_money += cst

                if cnt > 0 or not date_from:
                    data.append({
                        'fleet_number': v.internal_fleet_number,
                        'plate_number': v.plate_number,
                        'fuel_type': v.get_fuel_type_display(),
                        'refuel_count': cnt,
                        'total_liters': f"{lit:,.1f}",
                        'total_cost': f"{cst:,.2f}",
                    })
            summary = {
                'إجمالي لترات الوقود': f"{total_lit:,.1f} لتر",
                'إجمالي مصروفات الوقود': f"{total_money:,.2f} ر.س"
            }

        # 10. Odometer Report
        elif report_type == 'ODOMETER':
            title = "تقرير سجل وحركات العدادات"
            columns = [
                {'key': 'vehicle', 'label': 'المركبة'},
                {'key': 'reading', 'label': 'قراءة العداد (كم)'},
                {'key': 'date_time', 'label': 'التاريخ والوقت'},
                {'key': 'source', 'label': 'المصدر'},
                {'key': 'user', 'label': 'المسؤول'},
                {'key': 'notes', 'label': 'ملاحظات'},
            ]
            qs = OdometerReading.objects.all().select_related('vehicle', 'user')
            if date_from:
                qs = qs.filter(date_time__date__gte=date_from)
            if date_to:
                qs = qs.filter(date_time__date__lte=date_to)
            if vehicle_id:
                qs = qs.filter(vehicle_id=vehicle_id)

            for r in qs[:500]:
                data.append({
                    'vehicle': f"{r.vehicle.internal_fleet_number} ({r.vehicle.plate_number})",
                    'reading': f"{r.reading:,}",
                    'date_time': r.date_time.strftime('%Y-%m-%d %H:%M'),
                    'source': r.get_source_display(),
                    'user': r.user.full_name if r.user else 'نظام',
                    'notes': r.notes or '-',
                })
            summary = {'إجمالي القراءات المسجلة': len(data)}

        # 11. Driver Activity Report
        elif report_type == 'DRIVER_ACTIVITY':
            title = "تقرير نشاط وأداء السائقين"
            columns = [
                {'key': 'driver_name', 'label': 'اسم السائق'},
                {'key': 'national_id', 'label': 'الهوية الوطنية'},
                {'key': 'phone', 'label': 'رقم الجوال'},
                {'key': 'current_vehicle', 'label': 'المركبة المعينة حالياً'},
                {'key': 'total_trips', 'label': 'عدد الرحلات المنجزة'},
                {'key': 'total_distance', 'label': 'إجمالي المسافة (كم)'},
            ]
            qs = Driver.objects.filter(is_deleted=False)
            if driver_id:
                qs = qs.filter(id=driver_id)

            for d in qs:
                p_qs = d.work_permits.filter(status=WorkPermit.Status.COMPLETED)
                if date_from:
                    p_qs = p_qs.filter(departure_date__gte=date_from)
                if date_to:
                    p_qs = p_qs.filter(departure_date__lte=date_to)
                
                dist = p_qs.aggregate(total=Sum('distance_traveled'))['total'] or 0
                trips_cnt = p_qs.count()
                curr_v = d.current_vehicle

                data.append({
                    'driver_name': d.full_name,
                    'national_id': d.national_id,
                    'phone': d.phone,
                    'current_vehicle': f"{curr_v.internal_fleet_number} ({curr_v.plate_number})" if curr_v else 'غير معين',
                    'total_trips': trips_cnt,
                    'total_distance': f"{dist:,}",
                })
            summary = {'إجمالي السائقين': len(data)}

        # 12. Expiring Documents Report
        elif report_type == 'EXPIRING_DOCUMENTS':
            title = "تقرير الوثائق والتراخيص المنتهية أو القريبة من الانتهاء"
            columns = [
                {'key': 'vehicle', 'label': 'المركبة'},
                {'key': 'doc_type', 'label': 'نوع الوثيقة'},
                {'key': 'doc_number', 'label': 'رقم الوثيقة'},
                {'key': 'expiry_date', 'label': 'تاريخ الانتهاء'},
                {'key': 'days_left', 'label': 'الأيام المتبقية'},
                {'key': 'status', 'label': 'الحالة'},
            ]
            today = timezone.now().date()
            # Check Vehicle Documents
            docs = VehicleDocument.objects.all().select_related('vehicle')
            for doc in docs:
                days = (doc.expiry_date - today).days
                if days <= 60:
                    status_text = "منتهية" if days < 0 else "قاربت على الانتهاء"
                    data.append({
                        'vehicle': f"{doc.vehicle.internal_fleet_number} ({doc.vehicle.plate_number})",
                        'doc_type': doc.get_doc_type_display(),
                        'doc_number': doc.document_number or '-',
                        'expiry_date': str(doc.expiry_date),
                        'days_left': f"{days} يوم",
                        'status': status_text,
                    })

            # Check Vehicle Direct Fields (Registration & Insurance)
            vehicles = Vehicle.objects.filter(is_deleted=False)
            for v in vehicles:
                for label, exp in [('استمارة المركبة', v.registration_expiry_date), ('وثيقة التأمين', v.insurance_expiry_date), ('الفحص الدوري', v.periodic_inspection_expiry_date)]:
                    if exp:
                        days = (exp - today).days
                        if days <= 60:
                            data.append({
                                'vehicle': f"{v.internal_fleet_number} ({v.plate_number})",
                                'doc_type': label,
                                'doc_number': '-',
                                'expiry_date': str(exp),
                                'days_left': f"{days} يوم",
                                'status': "منتهية" if days < 0 else "قاربت على الانتهاء",
                            })

            summary = {'إجمالي الوثائق الحرجة': len(data)}

        # 13. Overdue Maintenance Report
        elif report_type == 'OVERDUE_MAINTENANCE':
            title = "تقرير الصيانات المتأخرة (حرجة)"
            columns = [
                {'key': 'vehicle', 'label': 'المركبة'},
                {'key': 'part_name', 'label': 'قطعة الغيار / الخدمة'},
                {'key': 'current_odometer', 'label': 'العداد الحالي'},
                {'key': 'next_km', 'label': 'العداد المستحق'},
                {'key': 'km_overdue', 'label': 'التجاوز بالكيلومتر'},
                {'key': 'next_date', 'label': 'التاريخ المستحق'},
                {'key': 'days_overdue', 'label': 'التجاوز بالأيام'},
            ]
            today = timezone.now().date()
            qs = VehiclePart.objects.filter(current_status=VehiclePart.HealthStatus.OVERDUE).select_related('vehicle', 'part')
            for p in qs:
                km_diff = p.vehicle.current_odometer - p.next_replacement_km
                days_diff = (today - p.next_replacement_date).days
                data.append({
                    'vehicle': f"{p.vehicle.internal_fleet_number} ({p.vehicle.plate_number})",
                    'part_name': p.part.name,
                    'current_odometer': f"{p.vehicle.current_odometer:,}",
                    'next_km': f"{p.next_replacement_km:,}",
                    'km_overdue': f"{max(km_diff, 0):,} كم",
                    'next_date': str(p.next_replacement_date),
                    'days_overdue': f"{max(days_diff, 0)} يوم",
                })
            summary = {'عدد الصيانات المتأخرة الحرجة': len(data)}

        # 14. Upcoming Maintenance Report
        elif report_type == 'UPCOMING_MAINTENANCE':
            title = "تقرير الصيانات القادمة والمستحقة قريباً"
            columns = [
                {'key': 'vehicle', 'label': 'المركبة'},
                {'key': 'part_name', 'label': 'قطعة الغيار / الخدمة'},
                {'key': 'status', 'label': 'الحالة'},
                {'key': 'current_odometer', 'label': 'العداد الحالي'},
                {'key': 'next_km', 'label': 'العداد القادم'},
                {'key': 'remaining_km', 'label': 'المتبقي (كم)'},
                {'key': 'next_date', 'label': 'التاريخ القادم'},
                {'key': 'remaining_days', 'label': 'المتبقي (أيام)'},
            ]
            today = timezone.now().date()
            qs = VehiclePart.objects.filter(current_status__in=[VehiclePart.HealthStatus.DUE_SOON, VehiclePart.HealthStatus.UPCOMING]).select_related('vehicle', 'part')
            for p in qs:
                rem_km = p.next_replacement_km - p.vehicle.current_odometer
                rem_days = (p.next_replacement_date - today).days
                data.append({
                    'vehicle': f"{p.vehicle.internal_fleet_number} ({p.vehicle.plate_number})",
                    'part_name': p.part.name,
                    'status': p.get_current_status_display(),
                    'current_odometer': f"{p.vehicle.current_odometer:,}",
                    'next_km': f"{p.next_replacement_km:,}",
                    'remaining_km': f"{rem_km:,}",
                    'next_date': str(p.next_replacement_date),
                    'remaining_days': f"{rem_days} يوم",
                })
            summary = {'عدد الصيانات المستحقة والقادمة': len(data)}

        else:
            title = "تقرير أسطول"

        return {
            'report_type': report_type,
            'title': title,
            'generated_at': timezone.now().strftime('%Y-%m-%d %H:%M'),
            'generated_by': user.full_name if user else 'نظام أسطول',
            'filters_used': filters,
            'columns': columns,
            'data': data,
            'summary': summary
        }

    @classmethod
    def export_excel(cls, report_type, filters, user=None):
        report = cls.generate_report_data(report_type, filters, user)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.views.sheetView[0].rightToLeft = True
        ws.title = report['title'][:30]

        # Colors & Fonts
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=16, bold=True, color="0F172A")
        meta_font = Font(name="Arial", size=10, italic=True, color="475569")
        data_font = Font(name="Arial", size=10)
        border_side = Side(border_style="thin", color="CBD5E1")
        thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

        # Title Row
        ws.merge_cells("A1:G1")
        ws["A1"] = f"نظام أسطول - {report['title']}"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # Metadata Row
        ws.merge_cells("A2:G2")
        ws["A2"] = f"تاريخ الاستخراج: {report['generated_at']} | المستخدم: {report['generated_by']}"
        ws["A2"].font = meta_font
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Table Header
        columns = report['columns']
        start_row = 4
        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=col['label'])
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[start_row].height = 25

        # Data Rows
        current_row = start_row + 1
        for row_data in report['data']:
            for col_idx, col in enumerate(columns, start=1):
                val = row_data.get(col['key'], '')
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = data_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
            ws.row_dimensions[current_row].height = 20
            current_row += 1

        # Summary Rows
        if report['summary']:
            current_row += 1
            for k, v in report['summary'].items():
                cell_label = ws.cell(row=current_row, column=1, value=k)
                cell_label.font = Font(name="Arial", size=11, bold=True)
                cell_val = ws.cell(row=current_row, column=2, value=str(v))
                cell_val.font = Font(name="Arial", size=11, bold=True)
                current_row += 1

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
